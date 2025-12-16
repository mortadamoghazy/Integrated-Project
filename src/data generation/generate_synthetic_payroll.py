import random
import numpy as np
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# ============================================================
# CONFIGURATION
# ============================================================

NUM_EMPLOYEES_SME = 250
NUM_EMPLOYEES_STARTUP = 30

NUM_MONTHS = 24
START_YEAR = 2023
START_MONTH = 1

# ============================================================
# MODELS
# ============================================================

MODEL_SME = {
    "Intern": {
        "pct": 0.05,
        "salary_range": (600, 900),
        "growth_range": (0.00, 0.00),
        "noise": 20,
        "bonus_prob": 0.00,
        "bonus_range": (0, 0),
    },
    "Junior": {
        "pct": 0.25,
        "salary_range": (1500, 2400),
        "growth_range": (0.003, 0.007),
        "noise": 30,
        "bonus_prob": 0.10,
        "bonus_range": (200, 800),
    },
    "Mid-Level": {
        "pct": 0.40,
        "salary_range": (2400, 3500),
        "growth_range": (0.0015, 0.005),
        "noise": 40,
        "bonus_prob": 0.20,
        "bonus_range": (300, 1200),
    },
    "Senior": {
        "pct": 0.20,
        "salary_range": (3500, 5500),
        "growth_range": (0.0005, 0.0025),
        "noise": 50,
        "bonus_prob": 0.25,
        "bonus_range": (400, 2500),
    },
    "Manager": {
        "pct": 0.10,
        "salary_range": (6000, 9000),
        "growth_range": (0.0003, 0.0020),
        "noise": 80,
        "bonus_prob": 0.40,
        "bonus_range": (1000, 7000),
    },
}

MODEL_STARTUP = {
    "Founder": {
        "pct": 0.10,
        "salary_range": (1500, 2500),
        "growth_range": (0.000, 0.002),
        "noise": 20,
        "bonus_prob": 0.02,
        "bonus_range": (100, 500),
    },
    "Engineer": {
        "pct": 0.60,
        "salary_range": (2200, 3500),
        "growth_range": (0.001, 0.004),
        "noise": 30,
        "bonus_prob": 0.05,
        "bonus_range": (100, 600),
    },
    "Junior": {
        "pct": 0.30,
        "salary_range": (1200, 1800),
        "growth_range": (0.000, 0.003),
        "noise": 20,
        "bonus_prob": 0.03,
        "bonus_range": (50, 300),
    },
}

# ============================================================
# HELPERS
# ============================================================

def get_role_model(model):
    return MODEL_STARTUP if model.upper() == "STARTUP" else MODEL_SME


def get_num_employees(model):
    return NUM_EMPLOYEES_STARTUP if model.upper() == "STARTUP" else NUM_EMPLOYEES_SME


def generate_employee_population(role_model, num_employees):
    employees = []
    employee_id = 1

    for role, cfg in role_model.items():
        count = int(num_employees * cfg["pct"])

        for _ in range(count):
            salary0 = random.uniform(*cfg["salary_range"])
            employees.append({
                "id": f"{employee_id:05d}",
                "role": role,
                "salary": salary0,
            })
            employee_id += 1

    random.shuffle(employees)
    return employees[:num_employees]


def compute_pas_rate(net_imp):
    if net_imp < 1500:
        return random.uniform(0.00, 0.00)
    elif net_imp < 2500:
        return random.uniform(0.03, 0.07)
    elif net_imp < 3500:
        return random.uniform(0.07, 0.11)
    elif net_imp < 5000:
        return random.uniform(0.11, 0.14)
    else:
        return random.uniform(0.14, 0.18)


def compute_cotisations(salary):
    sal_rate = random.uniform(0.18, 0.24)
    pat_rate = random.uniform(0.25, 0.42)
    return salary * sal_rate, salary * pat_rate


def month_name(year, month):
    return f"{year}-{month:02d}"


# ============================================================
# EXCEL FORMATTING
# ============================================================

def format_sheet(ws):
    # Auto column width
    for col in ws.columns:
        max_len = max(len(str(cell.value)) for cell in col if cell.value)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 3

    # Header formatting
    header_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    border_style = Side(border_style="thin", color="000000")
    border = Border(left=border_style, right=border_style, top=border_style, bottom=border_style)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    # Freeze header
    ws.freeze_panes = "A2"

    # Alternating rows
    alt_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if idx % 2 == 0:
            for cell in row:
                cell.fill = alt_fill

    # Number formatting
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")


# ============================================================
# MAIN GENERATOR
# ============================================================

def generate_payroll_excel(filename="synthetic_payroll.xlsx", model="SME"):
    role_model = get_role_model(model)
    num_employees = get_num_employees(model)

    employees = generate_employee_population(role_model, num_employees)

    wb = Workbook()
    wb.remove(wb.active)

    current_year, current_month = START_YEAR, START_MONTH

    for _ in range(NUM_MONTHS):
        sheet_title = month_name(current_year, current_month)
        ws = wb.create_sheet(sheet_title)

        headers = [
            "employee_id", "role", "salaire_brut", "cot_salariale",
            "cot_patronale", "net_imposable", "PAS", "net_paye",
            "avantages", "total_cost"
        ]
        ws.append(headers)

        for e in employees:
            cfg = role_model[e["role"]]

            # Salary update
            g = random.uniform(*cfg["growth_range"])
            e["salary"] *= 1 + g
            e["salary"] += np.random.normal(0, cfg["noise"])

            # Bonus
            if random.random() < cfg["bonus_prob"]:
                e["salary"] += random.uniform(*cfg["bonus_range"])

            salary = max(e["salary"], 0)

            cot_sal, cot_pat = compute_cotisations(salary)
            net_imp = salary - cot_sal
            pas = net_imp * compute_pas_rate(net_imp)
            net_paye = net_imp - pas
            avantages = random.uniform(0, 150)

            total_cost = salary + cot_pat + avantages

            ws.append([
                e["id"], e["role"], round(salary, 2),
                round(cot_sal, 2), round(cot_pat, 2),
                round(net_imp, 2), round(pas, 2),
                round(net_paye, 2), round(avantages, 2),
                round(total_cost, 2)
            ])

        format_sheet(ws)

        # Next month
        current_month += 1
        if current_month > 12:
            current_month = 1
            current_year += 1

    wb.save(filename)
    print(f"Generated {filename} using model {model}")
    

# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    generate_payroll_excel("synthetic_payroll_sme.xlsx", model="SME")
    generate_payroll_excel("synthetic_payroll_startup.xlsx", model="STARTUP")
